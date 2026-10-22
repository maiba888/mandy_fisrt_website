"""Factory order parsing, all-file validation, export and authenticated API tests."""
import http.client
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
from scripts import test_seo
import factory_orders as factory


def fixture_workbook(qty=5, info='01 * 5 PCS', price=1.2, offset=0, shade=None, missing_picture=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Order'
    for _ in range(offset):
        sheet.append(['Order metadata'])
    headers = ['NO', 'Brand', 'Item Name', 'Picture', 'Unit Price', 'QTY', 'Informations']
    row = [1, 'Brand', 'Lip gloss', None, price, qty, info]
    if shade is not None:
        headers.append('Shade/Type')
        row.append(shade)
    sheet.append(headers)
    sheet.append(row)
    if not missing_picture:
        image = io.BytesIO()
        Image.new('RGB', (30, 60), 'white').save(image, 'PNG')
        image.seek(0)
        sheet.add_image(XLImage(image), 'D' + str(offset + 2))
    sheet.append([None, None, None, None, None, 'TOTAL', '=BROKEN(1)'])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class FactoryTests(unittest.TestCase):
    def test_both_header_locations_and_shade_formats(self):
        files = [dict(filename='first.xlsx', data=fixture_workbook(qty=6, info='01 * 5 PCS\nwarm pink-2 * 1 PCS')),
                 dict(filename='second.xlsx', data=fixture_workbook(offset=3, shade='01'))]
        reports, errors, lines = factory.validate_factory_files(files)
        self.assertFalse(errors)
        self.assertEqual(sum(r['quantity'] for r in reports), 11)
        self.assertEqual([line['shade'] for line in lines], ['01', 'warm pink-2', '01'])
        self.assertEqual(len(factory.group_factory_lines(lines)), 2)

    def test_all_files_and_all_row_issues_reported(self):
        files = [dict(filename='broken.xlsx', data=b'not excel'),
                 dict(filename='invalid.xlsx', data=fixture_workbook(qty=0, info='wrong shade', price=-1, missing_picture=True)),
                 dict(filename='qty-mismatch.xlsx', data=fixture_workbook(qty=10))]
        reports, errors, _ = factory.validate_factory_files(files)
        self.assertEqual(len(reports), 3)
        self.assertEqual({error['filename'] for error in errors}, {f['filename'] for f in files})
        fields = {error['field'] for error in reports[1]['errors']}
        self.assertTrue({'qty', 'unit price', 'Informations', 'Picture'}.issubset(fields))
        self.assertTrue(any(error['row'] == 2 for error in reports[2]['errors']))

    def test_price_boundaries_and_source_filename_collision(self):
        files = [dict(filename='a.xlsx', data=fixture_workbook(price=1.2)),
                 dict(filename='b.xlsx', data=fixture_workbook(price=1.3)),
                 dict(filename='c.xlsx', data=fixture_workbook(price=1.20))]
        reports, errors, lines = factory.validate_factory_files(files)
        self.assertFalse(errors)
        groups = factory.group_factory_lines(lines)
        self.assertEqual(sorted(line['qty'] for line in groups), [5, 10])
        _, errors, _ = factory.validate_factory_files([files[0], files[0]])
        self.assertTrue(any('same filename' in error['reason'] for error in errors))

    def test_workbook_images_sources_fonts_and_owner(self):
        reports, errors, lines = factory.validate_factory_files([
            dict(filename='first.xlsx', data=fixture_workbook()),
            dict(filename='second.xlsx', data=fixture_workbook(qty=7, info='03 * 7 PCS'))])
        self.assertFalse(errors)
        with tempfile.TemporaryDirectory() as temp, patch.object(factory, 'BATCH_ROOT', Path(temp)):
            token, count = factory.store_factory_batch('admin', reports, lines)
            self.assertEqual(count, 2)
            with self.assertRaises(ValueError):
                factory.build_factory_workbook(token, 'other-admin')
            data = factory.build_factory_workbook(token, 'admin')
            book = load_workbook(io.BytesIO(data))
            self.assertEqual(book.sheetnames, ['Factory order', 'Sources'])
            self.assertEqual([cell.value for cell in book.active[1]], ['Product', 'Picture', 'Color / Type', 'QTY', 'Unit Price (USD)'])
            self.assertEqual([book.active.cell(r, 3).value for r in (2, 3)], ['01', '03'])
            self.assertEqual(book.active['C2'].number_format, '@')
            self.assertEqual(sum(book.active.cell(r, 4).value for r in (2, 3)), 12)
            self.assertEqual(len(book.active._images), 2)
            self.assertEqual([row[0].value for row in list(book['Sources'])[1:]], ['first.xlsx', 'second.xlsx'])
            self.assertEqual(book._named_styles['Normal'].font.name, 'Arial')
            self.assertTrue(all(cell.font.name == 'Arial' for sheet in book for row in sheet for cell in row))
            with ZipFile(io.BytesIO(data)) as archive:
                self.assertNotIn(b'<scheme', archive.read('xl/styles.xml'))
            book.close()


class FactoryApiTests(test_seo.SeoTests):
    def post(self, endpoint, data, content_type, admin=True):
        with patch.object(test_seo.server.CatalogHandler, 'current_user', return_value='admin' if admin else None):
            client = http.client.HTTPConnection('127.0.0.1', self.httpd.server_port, timeout=10)
            client.request('POST', endpoint, body=data, headers={'Content-Type': content_type})
            response = client.getresponse()
            result = response.status, response.read()
            client.close()
            return result

    def multipart(self, files):
        boundary = b'factory-test-boundary'
        body = b''
        for filename, data in files:
            body += b'--' + boundary + b'\r\nContent-Disposition: form-data; name="files"; filename="' + filename.encode() + b'"\r\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n' + data + b'\r\n'
        return body + b'--' + boundary + b'--\r\n', 'multipart/form-data; boundary=' + boundary.decode()

    def test_authentication_and_no_partial_merge(self):
        body, kind = self.multipart([('bad-a.xlsx', b'bad'), ('bad-b.xlsx', b'broken')])
        endpoint = '/api/admin/factory-orders/validate'
        self.assertEqual(self.post(endpoint, body, kind, admin=False)[0], 401)
        status, data = self.post(endpoint, body, kind)
        self.assertEqual(status, 422)
        result = json.loads(data)
        self.assertEqual(len(result['files']), 2)
        self.assertEqual({e['filename'] for e in result['errors']}, {'bad-a.xlsx', 'bad-b.xlsx'})
        self.assertNotIn('batch_token', result)
        self.assertEqual(self.post('/api/admin/factory-orders/merge', b'{}', 'application/json', admin=False)[0], 401)

    def test_valid_upload_merge_and_binary_integrity(self):
        data = fixture_workbook() + b'\r\n'
        body, kind = self.multipart([('customer-order.xlsx', data)])
        with tempfile.TemporaryDirectory() as temp, patch.object(factory, 'BATCH_ROOT', Path(temp)):
            status, response = self.post('/api/admin/factory-orders/validate', body, kind)
            self.assertEqual(status, 200, response)
            token = json.loads(response)['batch_token']
            status, output = self.post('/api/admin/factory-orders/merge', json.dumps({'batch_token': token}).encode(), 'application/json')
            self.assertEqual(status, 200)
            book = load_workbook(io.BytesIO(output))
            self.assertEqual(book.active['D2'].value, 5)
            book.close()


if __name__ == '__main__':
    unittest.main()

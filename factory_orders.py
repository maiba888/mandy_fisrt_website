"""Validate adjusted customer workbooks and prepare a separate factory order."""
import hashlib
import io
import json
import math
import re
import secrets
import shutil
import tempfile
import time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile, BadZipFile

MAX_FILES = 20
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_BATCH_BYTES = 180 * 1024 * 1024
BATCH_TTL = 2 * 60 * 60
BATCH_ROOT = Path(tempfile.gettempdir()) / 'luxe-factory-batches'


def text(value):
    return str(value if value is not None else '').strip()


def positive_number(value, integer=False):
    if isinstance(value, bool):
        raise ValueError('Must be a positive number.')
    try:
        result = Decimal(text(value))
    except InvalidOperation:
        raise ValueError('Must be a positive number, not text or an uncalculated formula.')
    if not result.is_finite() or result <= 0 or (integer and result != result.to_integral_value()):
        raise ValueError('Must be a positive whole number.' if integer else 'Must be greater than zero.')
    return int(result) if integer else result


def picture_bytes(image):
    from PIL import Image, ImageOps
    with Image.open(io.BytesIO(image._data())) as source:
        if source.width * source.height > 40_000_000:
            raise ValueError('Picture exceeds 40 megapixels.')
        picture = ImageOps.exif_transpose(source)
        picture.thumbnail((300, 600))
        picture = picture.convert('RGBA')
        background = Image.new('RGBA', picture.size, 'white')
        background.alpha_composite(picture)
        picture = background.convert('RGB')
        output = io.BytesIO()
        picture.save(output, 'JPEG', quality=92)
        return output.getvalue()


def validate_factory_files(files):
    from openpyxl import load_workbook
    reports, lines, filenames = [], [], set()
    for file in files:
        name, data = file['filename'], file['data']
        report = {'filename': name, 'errors': [], 'item_rows': 0, 'quantity': 0}
        reports.append(report)

        def error(reason, sheet='', row=None, field=''):
            report['errors'].append({'filename': name, 'sheet': sheet, 'row': row, 'field': field, 'reason': reason})

        if not name.lower().endswith('.xlsx'):
            error('Only .xlsx customer orders are supported.')
            continue
        if len(data) > MAX_FILE_BYTES:
            error('File exceeds the 50 MB limit.')
            continue
        if name in filenames:
            error('Another selected order uses the same filename. Rename the files so Sources can distinguish them.')
        filenames.add(name)
        workbook = None
        try:
            with ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 5000 or sum(entry.file_size for entry in archive.infolist()) > 250 * 1024 * 1024:
                    raise ValueError('Workbook is too large after decompression.')
            workbook = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
            found = False
            for sheet in workbook:
                if sheet.max_row > 5000 or sheet.max_column > 100:
                    error('Worksheet exceeds 5,000 rows or 100 columns.', sheet.title)
                    continue
                headers, header_row = None, None
                for row in sheet.iter_rows(max_row=min(25, sheet.max_row)):
                    labels = [text(cell.value).casefold() for cell in row]
                    if 'item name' in labels:
                        headers = {label: index for index, label in enumerate(labels) if label}
                        header_row = row[0].row
                        for label in set(labels) - {''}:
                            if labels.count(label) > 1:
                                error('Duplicate column heading.', sheet.title, header_row, label)
                        break
                if headers is None:
                    continue
                found = True
                required = ['brand', 'item name', 'picture', 'unit price', 'qty']
                missing = [field for field in required if field not in headers]
                if 'informations' not in headers and 'shade/type' not in headers:
                    missing.append('Informations or Shade/Type')
                if missing:
                    for field in missing:
                        error('Required column is missing.', sheet.title, header_row, field)
                    continue
                images = {}
                for image in sheet._images:
                    anchor = getattr(image.anchor, '_from', None)
                    if anchor is not None and anchor.col == headers['picture']:
                        images.setdefault(anchor.row + 1, image)
                for row in sheet.iter_rows(min_row=header_row + 1):
                    number = row[0].row
                    value = lambda field: row[headers[field]].value if field in headers else None
                    # Footer totals are not order detail, even when their formulas have errors.
                    footer_labels = {'total', 'total product cost', 'shipping cost', 'shipping country', 'shipping weight (kg)'}
                    if not text(value('brand')) and not text(value('item name')) and any(text(cell.value).casefold() in footer_labels for cell in row):
                        continue
                    if not any(text(value(field)) for field in ('brand', 'item name', 'unit price', 'qty')):
                        continue
                    start_errors = len(report['errors'])
                    brand, name_value = text(value('brand')), text(value('item name'))
                    for field, item in [('Brand', brand), ('Item Name', name_value)]:
                        if not item:
                            error('Required value is empty.', sheet.title, number, field)
                    price, quantity = None, None
                    for field, integer in [('unit price', False), ('qty', True)]:
                        try:
                            parsed = positive_number(value(field), integer)
                            if integer:
                                quantity = parsed
                            else:
                                price = parsed
                        except ValueError as exc:
                            error(str(exc), sheet.title, number, field)
                    shades = []
                    information = text(value('informations'))
                    if information:
                        for entry in information.splitlines():
                            if not entry.strip():
                                continue
                            match = re.fullmatch(r'(.+?)\s*\*\s*(\d+)\s*PCS\s*', entry.strip(), re.I)
                            if not match or int(match.group(2)) <= 0:
                                error('Invalid shade quantity: ' + entry.strip() + '. Use Shade * 5 PCS.', sheet.title, number, 'Informations')
                            else:
                                shades.append((match.group(1).strip(), int(match.group(2))))
                        if quantity is not None and sum(q for _, q in shades) != quantity:
                            error('Shade quantities total ' + str(sum(q for _, q in shades)) + ', but QTY is ' + str(quantity) + '.', sheet.title, number, 'Informations / QTY')
                        listed = [s.strip() for s in text(value('shade/type')).splitlines() if s.strip()]
                        if listed and set(listed) != {s for s, _ in shades}:
                            error('Shade/Type does not match the shades in Informations.', sheet.title, number, 'Shade/Type')
                    else:
                        listed = [s.strip() for s in text(value('shade/type')).splitlines() if s.strip()]
                        if len(listed) == 1 and quantity is not None:
                            shades = [(listed[0], quantity)]
                        else:
                            error('Provide Informations with quantities for every shade, or one Shade/Type with QTY.', sheet.title, number, 'Informations')
                    picture = None
                    if number not in images:
                        error('No embedded picture in the Picture column for this row.', sheet.title, number, 'Picture')
                    else:
                        try:
                            picture = picture_bytes(images[number])
                        except Exception:
                            error('Embedded picture cannot be read. Insert a valid JPG or PNG picture.', sheet.title, number, 'Picture')
                    if len(report['errors']) == start_errors:
                        report['item_rows'] += 1
                        report['quantity'] += quantity
                        lines.extend({'brand': brand, 'name': name_value, 'shade': shade, 'qty': qty,
                                      'price': str(price), 'picture': picture} for shade, qty in shades)
            if not found:
                error('Order header was not found within the first 25 rows of any worksheet.')
            if not report['item_rows'] and not report['errors']:
                error('Workbook has no order items.')
        except BadZipFile:
            error('Not a readable Excel .xlsx file. Save it again as .xlsx and re-upload.')
        except (OSError, ValueError, KeyError) as exc:
            error('Cannot read workbook: ' + str(exc))
        except Exception:
            error('Workbook is damaged or unsupported. Save it again as .xlsx and re-upload.')
        finally:
            if workbook is not None:
                workbook.close()
    errors = [error for report in reports for error in report['errors']]
    return reports, errors, lines


def group_factory_lines(lines):
    groups = {}
    for line in lines:
        key = (line['brand'], line['name'], line['shade'], Decimal(line['price']))
        if key not in groups:
            groups[key] = dict(line)
        else:
            groups[key]['qty'] += line['qty']
    return sorted(groups.values(), key=lambda line: (line['brand'].casefold(), line['name'].casefold(), line['shade'].casefold(), Decimal(line['price'])))


def store_factory_batch(owner, reports, lines):
    BATCH_ROOT.mkdir(mode=0o700, parents=True, exist_ok=True)
    root = BATCH_ROOT.resolve()
    for directory in root.iterdir():
        if directory.is_dir() and not directory.is_symlink() and re.fullmatch('[a-f0-9]{32}', directory.name) and time.time() - directory.stat().st_mtime > BATCH_TTL:
            if directory.resolve().parent == root:
                shutil.rmtree(directory)
    token = secrets.token_hex(16)
    directory = root / token
    directory.mkdir(mode=0o700)
    items = []
    for line in group_factory_lines(lines):
        image = line['picture']
        filename = hashlib.sha256(image).hexdigest() + '.jpg'
        target = directory / filename
        if not target.exists():
            target.write_bytes(image)
        items.append({**line, 'picture': filename})
    payload = {'owner': owner, 'created_at': time.time(), 'files': [r['filename'] for r in reports], 'items': items}
    (directory / 'manifest.json').write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
    return token, len(items)


def build_factory_workbook(token, owner):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill
    if not re.fullmatch('[a-f0-9]{32}', str(token)):
        raise ValueError('Upload and validate the customer orders first.')
    directory = BATCH_ROOT / token
    manifest = directory / 'manifest.json'
    if not manifest.is_file():
        raise ValueError('Upload expired. Please upload and validate the orders again.')
    payload = json.loads(manifest.read_text(encoding='utf-8'))
    if payload['owner'] != owner or time.time() - payload['created_at'] > BATCH_TTL:
        raise ValueError('Upload expired. Please upload and validate the orders again.')
    book = Workbook()
    default = Font(name='Arial', size=11, family=2)
    book._fonts[0] = default
    book._named_styles['Normal'].font = default
    sheet = book.active
    sheet.title = 'Factory order'
    sheet.append(['Product', 'Picture', 'Color / Type', 'QTY', 'Unit Price (USD)'])
    for number, item in enumerate(payload['items'], 2):
        sheet.append([item['brand'] + '\n' + item['name'], '', item['shade'], item['qty'], float(Decimal(item['price']))])
        # Imported text is literal text, never an Excel formula.
        sheet.cell(number, 1).data_type = 's'
        sheet.cell(number, 3).data_type = 's'
        sheet.cell(number, 3).number_format = '@'
        image = Image(str(directory / item['picture']))
        scale = min(119 / image.width, 400 / image.height)
        image.height *= scale
        image.width *= scale
        sheet.add_image(image, 'B' + str(number))
        text_height = 14 * (math.ceil(len(item['brand']) / 45) + math.ceil(len(item['name']) / 45)) + 8
        sheet.row_dimensions[number].height = max(50, image.height * .75 + 8, text_height)
        sheet.cell(number, 4).number_format = '0'
        sheet.cell(number, 5).number_format = '$0.00'
    for column, width in [('A', 52), ('B', 19), ('C', 28), ('D', 12), ('E', 20)]:
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sources = book.create_sheet('Sources')
    sources.append(['Source order file'])
    for filename in payload['files']:
        sources.append([filename])
        sources.cell(sources.max_row, 1).data_type = 's'
        sources.cell(sources.max_row, 1).number_format = '@'
    sources.column_dimensions['A'].width = 65
    sources.freeze_panes = 'A2'
    for tab in book:
        tab.sheet_view.showGridLines = False
        tab.row_dimensions[1].height = 28
        for row in tab:
            for cell in row:
                cell.font = default
                cell.alignment = Alignment(vertical='center', wrap_text=True)
        for cell in tab[1]:
            cell.font = Font(name='Arial', family=2, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F1A17')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
